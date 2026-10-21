"""Worker-only streaming ingestion + Spark artifacts; web reads bounded SQLite pages."""
import csv
from contextlib import closing, contextmanager
from datetime import date, datetime, time
import json
from pathlib import Path
import shutil
import sqlite3
from tempfile import TemporaryDirectory
import zipfile

from botocore.exceptions import BotoCoreError, ClientError
from celery.exceptions import SoftTimeLimitExceeded
from django.conf import settings

from .s3 import _translate_error, make_client
from .credentials import CredentialsUnavailable


class DatasetProblem(Exception):
    def __init__(self, code, message):
        self.code, self.message = code, message
        super().__init__(message)


class DatasetCancelled(Exception):
    pass


def artifact_directory(dataset_id):
    # UUIDs from model/API only; S3 keys and worksheet names never become paths.
    from uuid import UUID
    return settings.DATASET_ROOT / str(UUID(str(dataset_id)))


def _cell(value):
    if value is None or value == '':
        return None
    if isinstance(value, (date, datetime, time)):
        value = value.isoformat()
    else:
        value = str(value)
    if len(value) > settings.MAX_CELL_CHARACTERS:
        raise DatasetProblem('CELL_TOO_LARGE', 'A cell exceeds the 65,536-character import limit.')
    return value


def _excel_value(value, number_format):
    # Excel stores dates as datetimes even when the cell is date-only. Retain
    # genuine timestamps (including midnight timestamps) and time-only cells.
    from openpyxl.styles.numbers import is_datetime
    if (isinstance(value, datetime) and value.time() == time(0, 0)
            and is_datetime(number_format) == 'date'):
        return value.date()
    return value


def _columns(header):
    if not header or len(header) > settings.MAX_DATASET_COLUMNS:
        raise DatasetProblem('INVALID_HEADER', 'Use a header row with between 1 and 256 columns.')
    columns, used = [], set()
    for index, value in enumerate(header):
        name = str(value).strip() if value is not None else ''
        name = name or f'Column {index + 1}'
        if len(name) > 256:
            raise DatasetProblem('INVALID_HEADER', 'Column names must be at most 256 characters.')
        candidate, suffix = name, 2
        while candidate in used:
            candidate = f'{name} ({suffix})'
            suffix += 1
        used.add(candidate)
        columns.append({'key': f'c{index}', 'name': candidate, 'type': 'string'})
    return columns


def normalize_rows(rows, target, report):
    """Preserve order and leading zeros. Blank cells -> null; skip empty rows."""
    rows = iter(rows)
    try:
        header = next(rows)
    except StopIteration:
        raise DatasetProblem('EMPTY_FILE', 'The file has no header row.') from None
    columns = _columns(header)
    count = 0
    with Path(target).open('w', encoding='utf-8') as output:
        for source_row, row in enumerate(rows, start=2):
            if not row or all(value is None or value == '' for value in row):
                continue
            if len(row) != len(columns):
                raise DatasetProblem('ROW_WIDTH', f'Row {source_row} does not match the header width.')
            count += 1
            record = {'_row_id': count}
            record.update({column['key']: _cell(value) for column, value in zip(columns, row)})
            output.write(json.dumps(record, ensure_ascii=False) + '\n')
            if count % 1000 == 0:
                report(f'Parsing source rows ({count:,} read)', 35)
    report(f'Parsed {count:,} rows', 45)
    return columns, count


def parse_source(source, suffix, sheet_name, target, report):
    if suffix == '.csv':
        if sheet_name:
            raise DatasetProblem('CSV_SHEET', 'Worksheet selection is only available for Excel files.')
        try:
            csv.field_size_limit(settings.MAX_CELL_CHARACTERS)
            with Path(source).open(encoding='utf-8-sig', newline='') as stream:
                columns, count = normalize_rows(csv.reader(stream, strict=True), target, report)
            return columns, count, {'sheet_names': [], 'sheet_name': ''}
        except (UnicodeError, csv.Error):
            raise DatasetProblem('CSV_INVALID', 'Cannot read this CSV. Use UTF-8 encoding, comma delimiters and valid quoting.') from None
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        workbook = None
        try:
            with zipfile.ZipFile(source) as archive:
                entries = archive.infolist()
                if len(entries) > 10000 or sum(entry.file_size for entry in entries) > settings.MAX_EXCEL_EXPANDED_BYTES:
                    raise DatasetProblem('EXCEL_TOO_LARGE', 'The expanded Excel workbook exceeds the import limit.')
            workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
            names = workbook.sheetnames
            chosen = sheet_name or names[0]
            if chosen not in names:
                raise DatasetProblem('SHEET_NOT_FOUND', 'The worksheet name does not exist. Check its exact spelling.')
            sheet = workbook[chosen]
            # Do not trust a producer's cached dimensions: streaming iterates actual cells.
            sheet.reset_dimensions()
            # A trailing blank header is omitted from the first streamed row.
            # Scan actual populated cells, rather than guessing width from that
            # row or trusting potentially inflated producer dimensions.
            width = 0
            for row_number, row in enumerate(sheet.iter_rows(), 1):
                for index, cell in enumerate(row, 1):
                    if cell.value is not None and cell.value != '':
                        width = max(width, index)
                        if width > settings.MAX_DATASET_COLUMNS:
                            raise DatasetProblem('INVALID_HEADER', 'Use a header row with between 1 and 256 columns.')
                if row_number % 1000 == 0:
                    report(f'Parsing source rows (worksheet width scan: {row_number:,})', 35)
            report('Worksheet width checked', 35)
            def excel_rows():
                for row in sheet.iter_rows():
                    values = []
                    for cell in row:
                        if cell.data_type == 'f':
                            raise DatasetProblem('EXCEL_FORMULA', 'This worksheet contains formulas. Export a values-only copy before importing.')
                        values.append(_excel_value(cell.value, cell.number_format))
                    if len(values) < width:
                        values.extend([None] * (width - len(values)))
                    elif len(values) > width and all(value is None or value == '' for value in values[width:]):
                        values = values[:width]
                    yield values
            columns, count = normalize_rows(excel_rows(), target, report)
            return columns, count, {'sheet_names': names, 'sheet_name': chosen}
        except (DatasetProblem, DatasetCancelled, SoftTimeLimitExceeded):
            raise
        except Exception:
            raise DatasetProblem('EXCEL_INVALID', 'Cannot read this XLSX workbook. Check that it is valid and not password-protected.') from None
        finally:
            if workbook:
                workbook.close()
    if suffix == '.xls':
        import xlrd
        workbook = None
        try:
            workbook = xlrd.open_workbook(source, on_demand=True, formatting_info=True)
            names = workbook.sheet_names()
            chosen = sheet_name or names[0]
            if chosen not in names:
                raise DatasetProblem('SHEET_NOT_FOUND', 'The worksheet name does not exist. Check its exact spelling.')
            sheet = workbook.sheet_by_name(chosen)
            def legacy_rows():
                for index in range(sheet.nrows):
                    values = []
                    for cell in sheet.row(index):
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(cell.value, workbook.datemode)
                            fmt = workbook.format_map[workbook.xf_list[cell.xf_index].format_key].format_str
                            value = _excel_value(value, fmt)
                        elif cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                            value = None
                        elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            value = bool(cell.value)
                        elif cell.ctype == xlrd.XL_CELL_ERROR:
                            value = xlrd.error_text_from_code.get(cell.value, '#ERROR')
                        else:
                            value = cell.value
                        values.append(value)
                    yield values
            columns, count = normalize_rows(legacy_rows(), target, report)
            return columns, count, {'sheet_names': names, 'sheet_name': chosen}
        except (DatasetProblem, DatasetCancelled, SoftTimeLimitExceeded):
            raise
        except Exception:
            raise DatasetProblem('EXCEL_INVALID', 'Cannot read this XLS workbook. Check that it is valid and not password-protected.') from None
        finally:
            if workbook:
                workbook.release_resources()
    raise DatasetProblem('UNSUPPORTED_FILE', 'Choose a CSV, XLSX or XLS file.')


def download_source(connection, key, target, report):
    suffix = Path(key).suffix.lower()
    maximum = settings.MAX_SOURCE_BYTES if suffix == '.csv' else min(settings.MAX_SOURCE_BYTES, settings.MAX_EXCEL_BYTES)
    try:
        response = make_client(connection).get_object(Bucket=connection.bucket, Key=key)
        body = response['Body']
        try:
            size = response['ContentLength']
            if size > maximum:
                raise DatasetProblem('SOURCE_TOO_LARGE', f'This source exceeds the {maximum // 1048576} MB import limit.')
            total = 0
            with Path(target).open('wb') as stream:
                while chunk := body.read(512 * 1024):
                    total += len(chunk)
                    if total > maximum:
                        raise DatasetProblem('SOURCE_TOO_LARGE', 'This source exceeds the import size limit.')
                    stream.write(chunk)
                    report('Downloading source from S3', 15 + min(15, int(15 * total / max(size, 1))))
            if total != size:
                raise DatasetProblem('DOWNLOAD_INCOMPLETE', 'The S3 download was incomplete. Please retry.')
            return {'source_bytes': total, 'source_etag': response.get('ETag', ''),
                    'source_version': response.get('VersionId', '')}
        finally:
            body.close()
    except (BotoCoreError, ClientError) as exc:
        if isinstance(exc, ClientError) and exc.response.get('Error', {}).get('Code') == 'NoSuchKey':
            raise DatasetProblem('SOURCE_NOT_FOUND', 'The selected S3 file no longer exists. Refresh the file list.') from None
        raise _translate_error(exc) from None


def create_spark(job_id):
    from pyspark.sql import SparkSession
    spark = (SparkSession.builder.master(settings.SPARK_MASTER)
             .appName(f'Patternworks-{job_id}')
             .config('spark.ui.enabled', 'false')
             .config('spark.driver.bindAddress', '127.0.0.1')
             .config('spark.driver.host', '127.0.0.1')
             .config('spark.driver.memory', '1g')
             .config('spark.sql.shuffle.partitions', str(settings.SPARK_PARTITIONS))
             .config('spark.sql.session.timeZone', 'UTC')
             .config('spark.sql.debug.maxToStringFields', '0')
             .getOrCreate())
    spark.sparkContext.setLogLevel('OFF')
    spark.sparkContext.setJobGroup(str(job_id), 'Dataset preparation', interruptOnCancel=True)
    return spark


def build_page_index(json_directory, database_path, columns, report):
    from .content_identity import row_digest, content_digest
    keys = [column['key'] for column in columns]
    count, batch, batch_bytes = 0, [], 0
    with closing(sqlite3.connect(database_path)) as database, database:
        database.execute('CREATE TABLE rows (row_id INTEGER PRIMARY KEY, cells TEXT NOT NULL)')
        database.execute('CREATE TABLE full_samples (row_id INTEGER PRIMARY KEY, cells TEXT NOT NULL)')
        database.execute('CREATE TEMP TABLE content_rows (row_id INTEGER PRIMARY KEY, digest BLOB NOT NULL)')
        for part in sorted(Path(json_directory).glob('part-*.json')):
            with part.open(encoding='utf-8') as stream:
                for line in stream:
                    record = json.loads(line)
                    cells = [record.get(key) for key in keys]
                    database.execute('INSERT INTO content_rows VALUES (?, ?)',
                                     (record['_row_id'], row_digest(record['_row_id'], cells)))
                    if record['_row_id'] <= 20:
                        database.execute('INSERT INTO full_samples VALUES (?, ?)',
                                         (record['_row_id'], json.dumps(cells, ensure_ascii=False)))
                    # Full values stay in Parquet. Bound browser payload sizes independently.
                    cells = [value[:settings.PREVIEW_CELL_CHARACTERS] if isinstance(value, str) else value for value in cells]
                    payload = json.dumps(cells, ensure_ascii=False)
                    batch.append((record['_row_id'], payload))
                    batch_bytes += len(payload.encode('utf-8'))
                    count += 1
                    if len(batch) >= 500 or batch_bytes >= 4 * 1048576:
                        database.executemany('INSERT INTO rows VALUES (?, ?)', batch)
                        batch.clear()
                        batch_bytes = 0
                        report(f'Indexing preview ({count:,} rows)', 85)
        if batch:
            database.executemany('INSERT INTO rows VALUES (?, ?)', batch)
        signature = content_digest(columns, (row[0] for row in database.execute(
            'SELECT digest FROM content_rows ORDER BY row_id')))
        database.execute('CREATE TABLE content_identity (digest TEXT NOT NULL)')
        database.execute('INSERT INTO content_identity VALUES (?)', (signature,))
    return count


@contextmanager
def column_samples(dataset, columns, report):
    """Stream up to 20 complete rows, projecting only selected columns.

    New SQLite sample indexes and legacy Parquet share the same bounded path.
    Always close the cursor/Spark context, including cancellation and validation
    failures. A multi-column preview starts at most one legacy Spark context.
    """
    positions = {entry['key']: i for i, entry in enumerate(dataset.columns)}
    if not columns or len(set(columns)) != len(columns) or any(key not in positions for key in columns):
        raise DatasetProblem('COLUMN_UNAVAILABLE', 'Choose existing, unique target columns.')
    database = artifact_directory(dataset.pk) / 'rows.sqlite3'
    if database.is_file():
        with closing(sqlite3.connect(database.as_uri() + '?mode=ro', uri=True)) as connection:
            try:
                available = connection.execute("SELECT 1 FROM sqlite_master WHERE name='full_samples' AND type='table'").fetchone()
                if available:
                    projections = ', '.join('json_extract(cells, ?)' for _ in columns)
                    cursor = connection.execute(
                        f'SELECT row_id, {projections} FROM full_samples ORDER BY row_id LIMIT 20',
                        [f'$[{positions[key]}]' for key in columns])
                    yield ((row[0], list(row[1:])) for row in cursor)
                    return
            except sqlite3.Error:
                raise DatasetProblem('ARTIFACT_UNAVAILABLE', 'Cannot read the full-value sample. Import the file again.') from None
    parquet = artifact_directory(dataset.pk) / 'parquet'
    if not parquet.is_dir():
        raise DatasetProblem('ARTIFACT_UNAVAILABLE', 'The full-value source is unavailable. Import the file again.')
    report('Reading complete values from the source dataset', 80)
    job = getattr(report, 'job', None)
    spark = create_spark(job.pk if job else dataset.pk)
    try:
        if hasattr(report, 'bind_spark'):
            report.bind_spark(spark)
        from pyspark.sql import functions as f
        rows = (spark.read.parquet(str(parquet)).where(f.col('_row_id') <= 20)
                .select('_row_id', *columns).orderBy('_row_id').limit(20).toLocalIterator())
        yield ((row[0], list(row[1:])) for row in rows)
    finally:
        spark.stop()


def read_column_sample(dataset, column, report):
    """Compatibility wrapper for a single complete selected-column sample."""
    with column_samples(dataset, [column], report) as rows:
        return [(number, cells[0]) for number, cells in rows]


def _persist_dataframe(frame, dataset, staging, report):
    report('Writing partitioned Spark Parquet', 60)
    parquet = staging / 'parquet'
    frame.repartition(settings.SPARK_PARTITIONS).write.mode('error').parquet(str(parquet))
    report('Building paginated preview', 75)
    export = staging / 'preview-json'
    frame.sparkSession.read.parquet(str(parquet)).write.json(str(export))
    count = build_page_index(export, staging / 'rows.sqlite3', dataset.columns, report)
    with closing(sqlite3.connect(staging / 'rows.sqlite3')) as index:
        content_hash = index.execute('SELECT digest FROM content_identity').fetchone()[0]
    shutil.rmtree(export)
    return {'row_count': count, 'content_hash': content_hash, 'partition_count': len(list(parquet.glob('part-*.parquet'))),
            'spark_version': frame.sparkSession.version, 'spark_master': settings.SPARK_MASTER,
            'preview_cell_limit': settings.PREVIEW_CELL_CHARACTERS}


def prepare_dataset(job, report):
    dataset = job.dataset
    root = settings.DATASET_ROOT
    root.mkdir(parents=True, exist_ok=True)
    spark = None
    cached_frame = None
    try:
        if job.kind == 'APPLY_TRANSFORM':
            from .transform_store import prepare
            prepare(job.transformation, report)
        with TemporaryDirectory(prefix=f'.import-{job.pk}-{job.attempt_token}-', dir=root) as temporary:
            work = Path(temporary)
            staging = work / 'artifact'
            staging.mkdir()
            metadata = {}
            if job.kind == 'IMPORT_DATASET':
                if not job.connection:
                    raise CredentialsUnavailable()
                suffix = Path(dataset.source_key).suffix.lower()
                source = work / f'source{suffix}'
                metadata = download_source(job.connection, dataset.source_key, source, report)
                normalized = work / 'normalized.jsonl'
                columns, expected, workbook_info = parse_source(source, suffix, dataset.sheet_name, normalized, report)
                dataset.columns = columns
                dataset.sheet_name = workbook_info['sheet_name']
                metadata.update(workbook_info)
                report('Starting Spark DataFrame', 50)
                spark = create_spark(job.pk)
                if hasattr(report, 'bind_spark'):
                    report.bind_spark(spark)
                from pyspark.sql.types import LongType, StringType, StructField, StructType
                schema = StructType([StructField('_row_id', LongType(), False)] +
                                    [StructField(column['key'], StringType(), True) for column in columns])
                frame = spark.read.schema(schema).json(str(normalized))
            else:
                parent = dataset.parent
                source_path = artifact_directory(parent.pk) / 'parquet'
                if not source_path.is_dir():
                    raise DatasetProblem('ARTIFACT_UNAVAILABLE', 'The source dataset artifact is unavailable. Import it again.')
                dataset.columns = parent.columns
                expected = parent.row_count
                report('Starting Spark transformation', 40)
                spark = create_spark(job.pk)
                if hasattr(report, 'bind_spark'):
                    report.bind_spark(spark)
                from pyspark.sql import functions as functions
                frame = spark.read.parquet(str(source_path))
                if job.kind == 'APPLY_TRANSFORM':
                    from .transformations import apply_transformation
                    frame, metadata = apply_transformation(frame, dataset, job.transformation, report)
                elif job.kind == 'APPLY_RULE':
                    from functools import partial
                    from pyspark import StorageLevel
                    from pyspark.sql.types import ArrayType, BooleanType, LongType, StringType, StructField, StructType
                    from .regex_engine import RuleValidationError, apply_cells, validate_pattern
                    rule = job.rule
                    if not rule or rule.status != 'READY' or rule.dataset_id != parent.pk:
                        raise DatasetProblem('RULE_UNAVAILABLE', 'The validated rule is unavailable for this source dataset.')
                    try:
                        validate_pattern(rule.pattern, rule.ignore_case)
                    except RuleValidationError as exc:
                        raise DatasetProblem(exc.code, exc.message) from None
                    columns = rule.target_columns
                    if (not columns or len(set(columns)) != len(columns)
                            or not set(columns).issubset({entry['key'] for entry in parent.columns})):
                        raise DatasetProblem('RULE_COLUMN_MISSING', 'The rule target columns are unavailable.')
                    report('Applying validated RE2 rule to selected columns in Spark', 50)
                    cell_schema = StructType([StructField('value', StringType()), StructField('matched', BooleanType()),
                        StructField('match_count', LongType()), StructField('changed', BooleanType()), StructField('problem', StringType())])
                    transform = functions.udf(partial(apply_cells, pattern=rule.pattern, replacement=rule.replacement,
                        ignore_case=rule.ignore_case, maximum=settings.MAX_CELL_CHARACTERS), ArrayType(cell_schema, containsNull=False))
                    cached_frame = frame.withColumn('_rule_results',
                        transform(functions.array(*[functions.col(key) for key in columns]))).persist(StorageLevel.DISK_ONLY)
                    results = functions.col('_rule_results')
                    aggregates = [
                        functions.sum(functions.exists(results, lambda cell: cell['matched']).cast('long')).alias('matched_rows'),
                        functions.sum(functions.exists(results, lambda cell: cell['changed']).cast('long')).alias('changed_rows')]
                    for index, key in enumerate(columns):
                        cell = results[index]
                        aggregates.extend([
                            functions.sum(cell['matched'].cast('long')).alias(f'{key}_matched_rows'),
                            functions.sum(cell['match_count']).alias(f'{key}_match_count'),
                            functions.sum(cell['changed'].cast('long')).alias(f'{key}_changed_rows'),
                            functions.max(cell['problem']).alias(f'{key}_problem')])
                    statistics = cached_frame.agg(*aggregates).first()
                    for key in columns:
                        if statistics[f'{key}_problem']:
                            raise DatasetProblem(statistics[f'{key}_problem'],
                                'A selected column exceeds the replacement limits. No result was published; simplify and regenerate.')
                    column_stats = {key: {name: int(statistics[f'{key}_{name}'] or 0)
                        for name in ('matched_rows', 'match_count', 'changed_rows')} for key in columns}
                    replacements = {key: results[index]['value'].alias(key) for index, key in enumerate(columns)}
                    frame = cached_frame.select(*[replacements.get(key, functions.col(key)) for key in frame.columns])
                    metadata = {'operation': 'regex_replace', 'rule_id': str(rule.pk), 'column': rule.column, 'columns': columns,
                        'parent_id': str(parent.pk), 'regex_engine': 'RE2', 'pattern': rule.pattern,
                        'matched_rows': int(statistics['matched_rows'] or 0),
                        'changed_rows': int(statistics['changed_rows'] or 0),
                        'match_count': sum(info['match_count'] for info in column_stats.values()),
                        'matched_cells': sum(info['matched_rows'] for info in column_stats.values()),
                        'changed_cells': sum(info['changed_rows'] for info in column_stats.values()),
                        'column_stats': column_stats}

                else:
                    column, find, replacement = (job.parameters[name] for name in ('column', 'find', 'replacement'))
                    # Existing literal replacement remains Java-safe through literal escaping.
                    pattern = '\\Q' + find.replace('\\E', '\\E\\\\E\\Q') + '\\E'
                    replacement = replacement.replace('\\', '\\\\').replace('$', '\\$')
                    frame = frame.withColumn(column, functions.regexp_replace(functions.col(column), pattern, replacement))
                    if frame.filter(functions.length(functions.col(column)) > settings.MAX_CELL_CHARACTERS).limit(1).count():
                        raise DatasetProblem('REPLACEMENT_TOO_LARGE', 'The replacement would exceed the cell-size limit. Use a shorter replacement.')
                    metadata = {'operation': 'literal_replace', 'column': column, 'parent_id': str(parent.pk)}
            info = _persist_dataframe(frame, dataset, staging, report)
            if info['row_count'] != expected:
                raise DatasetProblem('ROW_COUNT_MISMATCH', 'Row count changed during processing; the result was not published.')
            metadata.update(info)
            report('Publishing dataset', 95)
            from .execution import publish_artifact
            publish_artifact(job, staging)
            return {'columns': dataset.columns, 'sheet_name': dataset.sheet_name,
                    'row_count': info['row_count'], 'metadata': metadata}
    finally:
        if cached_frame is not None:
            cached_frame.unpersist()
        if spark:
            spark.stop()


def read_page(dataset, after=0, limit=25):
    database = artifact_directory(dataset.pk) / 'rows.sqlite3'
    if not database.is_file():
        raise DatasetProblem('ARTIFACT_UNAVAILABLE', 'The preview artifact is unavailable. Import the file again.')
    try:
        values, payload_bytes, more = [], 0, False
        with closing(sqlite3.connect(database.as_uri() + '?mode=ro', uri=True)) as connection:
            cursor = connection.execute('SELECT row_id, cells FROM rows WHERE row_id > ? ORDER BY row_id LIMIT ?',
                                        (after, limit + 1))
            for row in cursor:
                size = len(row[1].encode('utf-8'))
                if len(values) >= limit or (values and payload_bytes + size > 2 * 1048576):
                    more = True
                    break
                values.append(row)
                payload_bytes += size
        return {'dataset_id': str(dataset.pk), 'columns': dataset.columns, 'total_rows': dataset.row_count,
                'rows': [{'row_number': row_id, 'cells': json.loads(cells)} for row_id, cells in values],
                'next_cursor': values[-1][0] if more else None,
                'preview_cell_limit': settings.PREVIEW_CELL_CHARACTERS}
    except (sqlite3.Error, ValueError):
        raise DatasetProblem('ARTIFACT_UNAVAILABLE', 'Cannot read the preview artifact. Import the file again.') from None


def discard_artifact(dataset_id):
    shutil.rmtree(artifact_directory(dataset_id), ignore_errors=True)
