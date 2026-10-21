"""Regression cases discovered by the full requirements audit."""
from datetime import datetime, time, timedelta
import json
from pathlib import Path
import sqlite3
from unittest.mock import patch

from django.core.cache import cache
from django.utils import timezone
import pytest

from processing.models import Dataset, Job, RegexRule, S3Connection, LLMTransformation
from processing.services.credentials import encrypt_credentials
from processing.services.datasets import (artifact_directory, build_page_index, parse_source, read_page,
    read_column_sample, create_spark)
from processing.services.date_policy import expected_date
from processing.services.rules import generate_rule, RuleProblem
from processing.tasks import run_job

pytestmark = pytest.mark.django_db

def quiet(*args): pass

def indexed_source(settings, tmp_path, values):
    settings.DATASET_ROOT = tmp_path / 'artifacts'
    record = Dataset.objects.create(owner_session='audit', status='READY', source_key='edge.csv', source_format='csv',
        row_count=len(values), columns=[{'key':'c0','name':'Text','type':'string'}])
    target = artifact_directory(record.pk); target.mkdir(parents=True)
    parts = tmp_path / 'parts'; parts.mkdir()
    (parts/'part-000.json').write_text('\n'.join(json.dumps({'_row_id':i,'c0':v}) for i,v in enumerate(values,1)))
    build_page_index(parts, target/'rows.sqlite3', record.columns, quiet)
    return record

@pytest.mark.parametrize('value,pattern,expected_count', [
    ('a'*4096+'Z', 'a$', 0),
    ('x'*4096+' one@example.com', r'[a-z]+@[a-z]+\.[a-z]+', 1),
], ids=['anchor-after-display-boundary','match-after-display-boundary'])
def test_full_value_preview_is_correct_but_display_stays_bounded(settings,tmp_path,value,pattern,expected_count):
    source=indexed_source(settings,tmp_path,[value,None])
    rule=RegexRule.objects.create(owner_session='audit',dataset=source,column='c0',instruction='Find pattern')
    output={'supported':True,'pattern':pattern,'replacement':'MASK','ignore_case':False,'explanation':'Replace matches','positive_examples':['a' if pattern=='a$' else 'one@example.com'],'negative_examples':['no match']}
    cache.clear()
    with patch('processing.services.rules.request_rule', return_value=(output,{})):
        result=generate_rule(rule,quiet,replacement='MASK')['preview']
    assert result['matched_rows']==expected_count and result['full_values'] is True
    assert result['examples'][0]['match_count']==expected_count
    assert result['examples'][0]['before_truncated'] is True
    assert len(result['examples'][0]['before'])==4096
    assert result['examples'][1]['before'] is None
    assert read_page(source)['rows'][0]['cells'][0]==value[:4096]
    assert read_column_sample(source,'c0',quiet)==[(1,value),(2,None)]


def test_full_sample_detects_growth_outside_displayed_text(settings,tmp_path):
    source=indexed_source(settings,tmp_path,['x'*4096+'a'*61000])
    rule=RegexRule.objects.create(owner_session='audit',dataset=source,column='c0',instruction='Find a')
    output={'supported':True,'pattern':'a','replacement':'aa','ignore_case':False,'explanation':'Replace a','positive_examples':['a'],'negative_examples':['zzz']}
    cache.clear()
    with patch('processing.services.rules.request_rule',return_value=(output,{})),pytest.raises(RuleProblem) as exc:
        generate_rule(rule,quiet,replacement='aa')
    assert exc.value.code=='OUTPUT_TOO_LARGE'


def test_full_sample_is_first_twenty_in_source_order(settings,tmp_path):
    values=['v'+str(i) for i in range(25)]
    source=indexed_source(settings,tmp_path,values)
    assert read_column_sample(source,'c0',quiet)==list(enumerate(values[:20],1))


@pytest.mark.parametrize('epoch_name',['WINDOWS_EPOCH','MAC_EPOCH'])
def test_excel_date_only_and_timestamps_keep_their_meaning(tmp_path,epoch_name):
    from openpyxl import Workbook
    from openpyxl.utils import datetime as excel_datetime
    book=Workbook(); book.epoch=getattr(excel_datetime,epoch_name)
    sheet=book.active; sheet.append(['Date','Timestamp','Midnight timestamp','Time','Hidden time'])
    sheet.append([datetime(2024,2,29),datetime(2026,10,18,12,34,56),datetime(2026,10,18),time(12,30),datetime(2026,10,18,12,30)])
    sheet['A2'].number_format='yyyy-mm-dd'; sheet['B2'].number_format='yyyy-mm-dd hh:mm:ss'
    sheet['C2'].number_format='yyyy-mm-dd hh:mm:ss'; sheet['D2'].number_format='hh:mm:ss'
    sheet['E2'].number_format='yyyy-mm-dd'
    path=tmp_path/'dates.xlsx'; book.save(path)
    out=tmp_path/'out.jsonl'; parse_source(path,'.xlsx','',out,quiet)
    row=json.loads(out.read_text())
    assert row=={'_row_id':1,'c0':'2024-02-29','c1':'2026-10-18T12:34:56','c2':'2026-10-18T00:00:00','c3':'12:30:00','c4':'2026-10-18T12:30:00'}
    assert expected_date(row['c0'],'REJECT')==('2024-02-29','normalized')
    assert expected_date(row['c1'],'REJECT')==(None,'invalid')
    assert expected_date(row['c2'],'REJECT')==(None,'invalid')


@pytest.mark.spark
def test_spark_full_sample_legacy_fallback_and_native_excel_date_pipeline(settings,tmp_path):
    from openpyxl import Workbook
    settings.DATASET_ROOT=tmp_path/'artifacts'
    connection=S3Connection.objects.create(owner_session='audit',bucket='audit-fixtures',
        encrypted_credentials=encrypt_credentials({'access_key_id':'FAKE','secret_access_key':'FAKE'}),
        status='READY',expires_at=timezone.now()+timedelta(hours=1))
    book=Workbook(); sheet=book.active
    sheet.append(['Date','Notes']); sheet.append([datetime(2024,2,29),'x'*4100+' one@example.com'])
    sheet['A2'].number_format='yyyy-mm-dd'
    fixture=tmp_path/'dates.xlsx'; book.save(fixture)
    source=Dataset.objects.create(owner_session='audit',connection=connection,source_key='dates.xlsx',source_format='xlsx')
    imported=Job.objects.create(owner_session='audit',kind='IMPORT_DATASET',dataset=source,connection=connection)
    def download(conn,key,target,report):
        import shutil
        shutil.copyfile(fixture,target)
        return {'source_bytes':fixture.stat().st_size}
    with patch('processing.services.datasets.download_source',side_effect=download): run_job.run(str(imported.pk))
    imported.refresh_from_db(); source.refresh_from_db()
    assert imported.status=='SUCCESS'
    # Simulate the exact artifact schema used before this fix: projection must
    # recover complete values from Parquet, without altering the old index.
    with sqlite3.connect(artifact_directory(source.pk)/'rows.sqlite3') as db: db.execute('DROP TABLE full_samples')
    assert read_column_sample(source,'c1',quiet)==[(1,'x'*4100+' one@example.com')]
    with sqlite3.connect(artifact_directory(source.pk)/'rows.sqlite3') as db:
        assert not db.execute("SELECT 1 FROM sqlite_master WHERE name='full_samples'").fetchone()
    rule=RegexRule.objects.create(owner_session='audit',dataset=source,column='c1',instruction='Find emails')
    generated=Job.objects.create(owner_session='audit',kind='GENERATE_RULE',rule=rule,parameters={'replacement':'MASK'})
    response={'supported':True,'pattern':r'[a-z]+@[a-z]+\.[a-z]+','replacement':'MASK','ignore_case':False,'explanation':'Mask emails','positive_examples':['one@example.com'],'negative_examples':['no email']}
    with patch('processing.services.rules.request_rule',return_value=(response,{})): run_job.run(str(generated.pk))
    rule.refresh_from_db(); assert rule.preview['matched_rows']==1
    output=Dataset.objects.create(owner_session='audit',parent=source,source_key='dates.xlsx',source_format='xlsx')
    applied=Job.objects.create(owner_session='audit',kind='APPLY_RULE',rule=rule,dataset=output)
    run_job.run(str(applied.pk)); applied.refresh_from_db(); output.refresh_from_db()
    assert applied.status=='SUCCESS' and output.metadata['matched_rows']==1
    assert read_column_sample(output,'c1',quiet)==[(1,'x'*4100+' MASK')]
    plan=LLMTransformation.objects.create(owner_session='audit',dataset=output,kind='DATE',column='c0',instruction='',
        options={'date_order':'REJECT'},model='gpt-4.1-mini')
    preview=Job.objects.create(owner_session='audit',kind='PREVIEW_TRANSFORM',transformation=plan)
    def model(record,values):
        assert values==['2024-02-29']
        return [{'value':'2024-02-29','status':'normalized'}],{}
    with patch('processing.services.transformations.request_batch',side_effect=model): run_job.run(str(preview.pk))
    plan.refresh_from_db(); assert plan.preview['examples'][0]['after']=='2024-02-29'
    final=Dataset.objects.create(owner_session='audit',parent=output,source_key='dates.xlsx',source_format='xlsx')
    task=Job.objects.create(owner_session='audit',kind='APPLY_TRANSFORM',dataset=final,transformation=plan)
    with patch('processing.services.transformations.request_batch',side_effect=model): run_job.run(str(task.pk))
    task.refresh_from_db(); final.refresh_from_db()
    assert task.status=='SUCCESS' and read_page(final)['rows'][0]['cells'][-2:]==['2024-02-29','normalized']
    assert read_column_sample(source,'c1',quiet)==[(1,'x'*4100+' one@example.com')]

@pytest.mark.parametrize('header,values,names',[
    (['Name',None],['Alice','Useful'],['Name','Column 2']),
    (['Name',None,'Name'],['Alice','middle','other'],['Name','Column 2','Name (2)']),
])
def test_sparse_excel_headers_use_populated_sheet_width(tmp_path,header,values,names):
    from openpyxl import Workbook
    book=Workbook(); sheet=book.active; sheet.append(header); sheet.append(values)
    sheet.cell(row=3,column=40).number_format='0.00'  # formatted-only columns do not become data
    path=tmp_path/'sparse.xlsx'; book.save(path)
    out=tmp_path/'out.jsonl'
    columns,count,_=parse_source(path,'.xlsx','',out,quiet)
    assert count==1 and [x['name'] for x in columns]==names
    assert [json.loads(out.read_text())[f'c{i}'] for i in range(len(names))]==values


def test_excel_populated_overwide_sheet_rejected_before_normalization(tmp_path):
    from openpyxl import Workbook
    from processing.services.datasets import DatasetProblem
    book=Workbook(); sheet=book.active; sheet.append(['Name']); sheet.cell(row=2,column=257,value='data')
    path=tmp_path/'wide.xlsx'; book.save(path)
    with pytest.raises(DatasetProblem) as exc: parse_source(path,'.xlsx','',tmp_path/'out',quiet)
    assert exc.value.code=='INVALID_HEADER'


def test_cancelled_cleanup_advances_across_multiple_batches(settings,tmp_path):
    from processing.services.recovery import recover_jobs
    import uuid
    settings.DATASET_ROOT=tmp_path
    source=Dataset.objects.create(owner_session='audit',source_key='x.csv',source_format='csv',status='READY')
    original=artifact_directory(source.pk); original.mkdir(); (original/'keep').write_text('source')
    jobs=[]
    for i in range(105):
        output=Dataset.objects.create(owner_session='audit',parent=source,source_key='x.csv',source_format='csv',status='CANCELLED')
        target=artifact_directory(output.pk); target.mkdir(); (target/'partial').write_text('partial')
        jobs.append(Job.objects.create(owner_session='audit',kind='APPLY_RULE',dataset=output,status='CANCELLED',
            attempt_token=uuid.uuid4(),heartbeat_at=timezone.now()-timedelta(seconds=200),finished_at=timezone.now()-timedelta(seconds=200)))
    recover_jobs(); assert Job.objects.filter(cleanup_completed_at__isnull=False).count()==100
    recover_jobs(); assert Job.objects.filter(cleanup_completed_at__isnull=False).count()==105
    assert all(not artifact_directory(j.dataset_id).exists() for j in jobs)
    assert (original/'keep').read_text()=='source'
    assert Job.objects.filter(status='CANCELLED').count()==105


def test_cancelled_unclaimed_job_also_finishes_cleanup(settings,tmp_path):
    from processing.services.recovery import recover_jobs
    settings.DATASET_ROOT=tmp_path
    job=Job.objects.create(owner_session='audit',kind='DIAGNOSTIC',status='CANCELLED',
        finished_at=timezone.now()-timedelta(seconds=200))
    recover_jobs(); job.refresh_from_db(); assert job.cleanup_completed_at is not None

@pytest.mark.parametrize('pattern,positive,instruction',[
    (r'04\d{8}','0412345678','Find mobiles exactly as 04 followed by eight digits'),
    (r'ORD-[0-9]{6}','ORD-123456','Find ORD- followed by exactly six digits'),
])
def test_exact_identifier_validation_rejects_numeric_prefix_matches(pattern,positive,instruction):
    from processing.services.rules import validate_output
    response={'supported':True,'pattern':pattern,'replacement':'MASK','ignore_case':False,'explanation':'Mask',
              'positive_examples':[positive],'negative_examples':['unrelated']}
    with pytest.raises(RuleProblem) as exc: validate_output(response,instruction)
    assert exc.value.code=='RULE_EXAMPLES_FAILED'
    response['pattern']=r'\b'+pattern+r'\b'
    assert validate_output(response,instruction)==response


def test_generated_examples_are_checked_before_caching(settings,tmp_path):
    source=indexed_source(settings,tmp_path,['ORD-123456'])
    rule=RegexRule.objects.create(owner_session='audit',dataset=source,column='c0',instruction='Find order IDs')
    output={'supported':True,'pattern':'ORD-[0-9]{6}','replacement':'MASK','ignore_case':False,'explanation':'Mask',
            'positive_examples':['ORD-123456'],'negative_examples':['ORD-1234567']}
    cache.clear()
    from processing.services.rules import rule_cache_key
    with patch('processing.services.rules.request_rule',return_value=(output,{})),pytest.raises(RuleProblem) as exc:
        generate_rule(rule,quiet)
    assert exc.value.code=='RULE_EXAMPLES_FAILED'
    assert cache.get(rule_cache_key('audit',rule.instruction)) is None


def test_exact_wording_does_not_impose_fixed_numeric_length_on_emails():
    from processing.services.rules import validate_output
    response={'supported':True,'pattern':r'[a-z0-9]+@[a-z]+\.[a-z]+','replacement':'MASK',
        'ignore_case':False,'explanation':'Mask emails','positive_examples':['user123@example.com'],
        'negative_examples':['no email']}
    assert validate_output(response,'Match exactly valid email addresses in text.')==response
