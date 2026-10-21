from datetime import timedelta
import io
import json
from pathlib import Path
from unittest.mock import patch
import zipfile

from botocore.response import StreamingBody
from django.utils import timezone
import pytest
from rest_framework.test import APIClient

from processing.models import Dataset, Job, S3Connection
from processing.services.credentials import encrypt_credentials
from processing.services.datasets import (DatasetCancelled, DatasetProblem, artifact_directory, build_page_index,
    download_source, normalize_rows, parse_source, prepare_dataset, read_page)
from processing.tasks import run_job

pytestmark = pytest.mark.django_db
SAMPLES = Path(__file__).resolve().parents[2] / 'samples'


@pytest.fixture
def store(settings, tmp_path):
    settings.DATASET_ROOT = tmp_path / 'artifacts'
    settings.DATASET_ROOT.mkdir()
    return settings.DATASET_ROOT


@pytest.fixture
def browser(settings):
    settings.DEBUG = True
    client = APIClient(enforce_csrf_checks=True)
    session = client.get('/api/session/')
    client.credentials(HTTP_X_CSRFTOKEN=session.data['csrf_token'])
    return client


def connection(owner):
    return S3Connection.objects.create(owner_session=owner, bucket='test-data-bucket', region='ap-southeast-2',
        encrypted_credentials=encrypt_credentials({'access_key_id': 'TEST', 'secret_access_key': 'FAKE'}),
        status=S3Connection.Status.READY, expires_at=timezone.now() + timedelta(hours=1))


def dataset(owner='owner', **kwargs):
    return Dataset.objects.create(owner_session=owner, source_key='contacts.csv', source_format='csv', **kwargs)


def quiet(*args):
    pass


def test_csv_preserves_quotes_multiline_bom_nulls_and_leading_zero(tmp_path):
    source = tmp_path / 'data.csv'
    source.write_text('\ufeffID,Notes,Empty\n001,"a,b\nc",\n\n002,plain,x\n', encoding='utf-8')
    target = tmp_path / 'out.json'
    columns, count, info = parse_source(source, '.csv', '', target, quiet)
    rows = [json.loads(line) for line in target.read_text().splitlines()]
    assert count == 2
    assert [entry['name'] for entry in columns] == ['ID', 'Notes', 'Empty']
    assert rows[0] == {'_row_id': 1, 'c0': '001', 'c1': 'a,b\nc', 'c2': None}
    assert info['sheet_names'] == []


def test_duplicate_and_blank_headers_have_unique_names(tmp_path):
    columns, count = normalize_rows([['A', 'A', None, 'A (2)'], [1, 2, 3, 4]], tmp_path / 'out', quiet)
    assert [entry['name'] for entry in columns] == ['A', 'A (2)', 'Column 3', 'A (2) (2)']
    assert [entry['key'] for entry in columns] == ['c0', 'c1', 'c2', 'c3']
    assert count == 1


@pytest.mark.parametrize('content,code', [('', 'EMPTY_FILE'), ('A,B\n1\n', 'ROW_WIDTH'),
    ('A\n"unterminated', 'CSV_INVALID')])
def test_malformed_csv_fails_with_safe_error(tmp_path, content, code):
    source = tmp_path / 'data.csv'
    source.write_text(content)
    with pytest.raises(DatasetProblem) as caught:
        parse_source(source, '.csv', '', tmp_path / 'out', quiet)
    assert caught.value.code == code
    assert 'unterminated' not in caught.value.message


def test_header_only_file_is_a_valid_empty_dataset(tmp_path):
    columns, count = normalize_rows([['ID', 'Name']], tmp_path / 'out', quiet)
    assert count == 0 and len(columns) == 2


def test_cell_size_is_bounded(settings, tmp_path):
    settings.MAX_CELL_CHARACTERS = 4
    with pytest.raises(DatasetProblem, match='cell exceeds'):
        normalize_rows([['A'], ['12345']], tmp_path / 'out', quiet)


def test_csv_invalid_encoding_is_not_silently_replaced(tmp_path):
    source = tmp_path / 'bad.csv'
    source.write_bytes(b'A\n\xff')
    with pytest.raises(DatasetProblem) as caught:
        parse_source(source, '.csv', '', tmp_path / 'out', quiet)
    assert caught.value.code == 'CSV_INVALID'


def test_samples_csv_and_xlsx_are_identical(tmp_path):
    csv_target, excel_target = tmp_path / 'csv.json', tmp_path / 'excel.json'
    csv_columns, csv_count, _ = parse_source(SAMPLES / 'contacts.csv', '.csv', '', csv_target, quiet)
    xlsx_columns, xlsx_count, workbook = parse_source(SAMPLES / 'contacts.xlsx', '.xlsx', 'Contacts', excel_target, quiet)
    assert csv_count == xlsx_count == 8
    assert csv_columns == xlsx_columns
    assert csv_target.read_text() == excel_target.read_text()
    assert workbook == {'sheet_name': 'Contacts', 'sheet_names': ['Contacts']}


def test_xlsx_sheet_selection_dates_and_missing_trailing_cells(tmp_path):
    from datetime import datetime
    from openpyxl import Workbook
    source = tmp_path / 'book.xlsx'
    book = Workbook()
    book.active.title = 'Cover'
    book.active.append(['Cover only'])
    sheet = book.create_sheet('Data')
    sheet.append(['ID', 'Date', 'Missing'])
    sheet.append(['007', datetime(2026, 10, 17)])
    book.save(source)
    output = tmp_path / 'out.json'
    _, count, info = parse_source(source, '.xlsx', 'Data', output, quiet)
    assert count == 1 and info['sheet_name'] == 'Data'
    row = json.loads(output.read_text())
    assert row['c0'] == '007' and row['c1'].startswith('2026-10-17') and row['c2'] is None
    with pytest.raises(DatasetProblem) as caught:
        parse_source(source, '.xlsx', 'Not here', output, quiet)
    assert caught.value.code == 'SHEET_NOT_FOUND'


def test_xlsx_formulas_are_rejected_instead_of_returning_missing_cached_values(tmp_path):
    from openpyxl import Workbook
    source = tmp_path / 'formula.xlsx'
    book = Workbook()
    book.active.append(['Value'])
    book.active.append(['=1+1'])
    book.save(source)
    with pytest.raises(DatasetProblem) as caught:
        parse_source(source, '.xlsx', '', tmp_path / 'out', quiet)
    assert caught.value.code == 'EXCEL_FORMULA'


def test_xlsx_uncompressed_size_is_bounded(settings, tmp_path):
    settings.MAX_EXCEL_EXPANDED_BYTES = 100
    source = tmp_path / 'large.xlsx'
    with zipfile.ZipFile(source, 'w') as archive:
        archive.writestr('xl/large.xml', 'x' * 101)
    with pytest.raises(DatasetProblem) as caught:
        parse_source(source, '.xlsx', '', tmp_path / 'out', quiet)
    assert caught.value.code == 'EXCEL_TOO_LARGE'


def test_cancellation_during_excel_parse_is_preserved(tmp_path):
    def cancelled(*args):
        raise DatasetCancelled()
    with pytest.raises(DatasetCancelled):
        parse_source(SAMPLES / 'contacts.xlsx', '.xlsx', '', tmp_path / 'out', cancelled)


def test_download_streaming_checks_length_and_closes_body(settings, tmp_path):
    settings.MAX_SOURCE_BYTES = 10
    body = StreamingBody(io.BytesIO(b'a' * 11), 11)
    with patch('processing.services.datasets.make_client') as client:
        client.return_value.get_object.return_value = {'Body': body, 'ContentLength': 11}
        with pytest.raises(DatasetProblem) as caught:
            download_source(connection('owner'), 'data.csv', tmp_path / 'download', quiet)
    assert caught.value.code == 'SOURCE_TOO_LARGE'
    assert body._raw_stream.closed


def test_page_index_is_stable_bounded_and_preserves_nulls(store, tmp_path):
    record = dataset(status=Dataset.Status.READY)
    record.columns = [{'key': 'c0', 'name': 'Value', 'type': 'string'}]
    record.row_count = 3
    target = artifact_directory(record.pk)
    target.mkdir()
    parts = tmp_path / 'parts'
    parts.mkdir()
    (parts / 'part-00000.json').write_text('\n'.join(json.dumps({'_row_id': index, 'c0': value})
        for index, value in [(3, None), (1, 'first'), (2, 'x' * 5000)]))
    assert build_page_index(parts, target / 'rows.sqlite3', record.columns, quiet) == 3
    first = read_page(record, limit=2)
    second = read_page(record, after=first['next_cursor'], limit=2)
    assert [row['row_number'] for row in first['rows']] == [1, 2]
    assert len(first['rows'][1]['cells'][0]) == 4096
    assert second['rows'] == [{'row_number': 3, 'cells': [None]}]
    assert second['next_cursor'] is None


def test_preview_page_byte_budget_preserves_cursor_without_dropping_rows(store, tmp_path):
    record = dataset(status=Dataset.Status.READY)
    record.columns = [{'key': f'c{index}', 'name': f'Column {index}', 'type': 'string'} for index in range(200)]
    target = artifact_directory(record.pk)
    target.mkdir()
    parts = tmp_path / 'wide-parts'
    parts.mkdir()
    with (parts / 'part-00000.json').open('w') as stream:
        for row_id in range(1, 6):
            stream.write(json.dumps({'_row_id': row_id, **{entry['key']: 'x' * 4096 for entry in record.columns}}) + '\n')
    record.row_count = build_page_index(parts, target / 'rows.sqlite3', record.columns, quiet)
    rows, after = [], 0
    while True:
        page = read_page(record, after=after, limit=100)
        assert len(page['rows']) <= 2
        rows.extend(row['row_number'] for row in page['rows'])
        after = page['next_cursor']
        if after is None:
            break
    assert rows == [1, 2, 3, 4, 5]


@patch('processing.views.dispatch', return_value=True)
def test_import_returns_immediate_job_and_does_not_download(enqueue, browser):
    record = connection(browser.session.session_key)
    with patch('processing.services.datasets.download_source') as download:
        response = browser.post('/api/datasets/', {'connection_id': str(record.pk), 'key': 'contacts.xlsx'}, format='json')
    assert response.status_code == 202
    assert response.data['kind'] == 'IMPORT_DATASET'
    assert response.data['dataset_id'] and response.data['status'] == 'QUEUED'
    assert set(response.data['result']) == set()
    download.assert_not_called()
    enqueue.assert_called_once()


def test_import_requires_owned_ready_unexpired_connection(browser):
    record = connection('different-owner')
    data = {'connection_id': str(record.pk), 'key': 'contacts.csv'}
    assert browser.post('/api/datasets/', data, format='json').status_code == 404
    record.owner_session = browser.session.session_key
    record.expires_at = timezone.now() - timedelta(seconds=1)
    record.save()
    assert browser.post('/api/datasets/', data, format='json').status_code == 410
    record.expires_at = timezone.now() + timedelta(hours=1)
    record.status = S3Connection.Status.PENDING
    record.save()
    assert browser.post('/api/datasets/', data, format='json').status_code == 409
    assert Dataset.objects.count() == 0


def test_other_session_cannot_read_preview_metadata_or_replace(browser):
    record = dataset(status=Dataset.Status.READY)
    assert browser.get(f'/api/datasets/{record.pk}/').status_code == 404
    assert browser.get(f'/api/datasets/{record.pk}/rows/').status_code == 404
    assert browser.post(f'/api/datasets/{record.pk}/replace/', {}, format='json').status_code == 404
    assert browser.get('/api/datasets/').data['datasets'] == []


def test_reload_restores_only_owned_unexpired_connection_metadata(browser):
    owned = connection(browser.session.session_key)
    connection('different-owner')
    expired = connection(browser.session.session_key)
    expired.expires_at = timezone.now() - timedelta(seconds=1)
    expired.save()
    response = browser.get('/api/connections/')
    assert response.status_code == 200
    assert len(response.data['connections']) == 1
    entry = response.data['connections'][0]
    assert str(entry['id']) == str(owned.pk)
    assert set(entry) == {'id', 'bucket', 'region', 'expires_at'}
    assert response['Cache-Control'] == 'no-store'


@pytest.mark.parametrize('query', ['limit=0', 'limit=101', 'after=-1', 'after=not-a-number'])
def test_rows_reject_unbounded_or_invalid_page_requests(browser, query):
    record = dataset(browser.session.session_key, status=Dataset.Status.READY)
    assert browser.get(f'/api/datasets/{record.pk}/rows/?{query}').status_code == 400


def test_missing_artifact_is_reported_as_gone(browser, store):
    record = dataset(browser.session.session_key, status=Dataset.Status.READY)
    response = browser.get(f'/api/datasets/{record.pk}/rows/')
    assert response.status_code == 410
    assert 'ARTIFACT_UNAVAILABLE' == response.data['code']


def test_cancel_import_preserves_connection_credentials(browser, store):
    record = connection(browser.session.session_key)
    result = dataset(browser.session.session_key, connection=record)
    job = Job.objects.create(owner_session=browser.session.session_key, kind=Job.Kind.IMPORT_DATASET,
        connection=record, dataset=result)
    response = browser.post(f'/api/jobs/{job.pk}/cancel/', {}, format='json')
    assert response.data['status'] == Job.Status.CANCELLED
    result.refresh_from_db()
    record.refresh_from_db()
    assert result.status == Dataset.Status.CANCELLED
    assert record.status == S3Connection.Status.READY and record.encrypted_credentials


def test_parse_failure_does_not_destroy_reusable_s3_connection(store):
    record = connection('owner')
    output = dataset(connection=record)
    job = Job.objects.create(owner_session='owner', kind=Job.Kind.IMPORT_DATASET, connection=record, dataset=output)
    with patch('processing.tasks.prepare_dataset', side_effect=DatasetProblem('CSV_INVALID', 'Invalid CSV.')):
        run_job.run(str(job.pk))
    job.refresh_from_db()
    record.refresh_from_db()
    output.refresh_from_db()
    assert job.error_code == 'CSV_INVALID' and job.status == Job.Status.FAILED
    assert output.status == Dataset.Status.FAILED
    assert record.status == S3Connection.Status.READY and record.encrypted_credentials


def test_cancellation_after_artifact_creation_never_publishes_dataset(store):
    output = dataset()
    job = Job.objects.create(owner_session='owner', kind=Job.Kind.IMPORT_DATASET, dataset=output)
    target = artifact_directory(output.pk)
    target.mkdir()
    def finish_then_cancel(*args):
        Job.objects.filter(pk=job.pk).update(status=Job.Status.CANCELLED)
        return {'columns': [], 'row_count': 0, 'metadata': {}, 'sheet_name': ''}
    with patch('processing.tasks.prepare_dataset', side_effect=finish_then_cancel):
        run_job.run(str(job.pk))
    output.refresh_from_db()
    assert output.status != Dataset.Status.READY
    assert not target.exists()


@patch('processing.views.dispatch', return_value=True)
def test_replacement_creates_new_dataset_and_validates_column(enqueue, browser):
    source = dataset(browser.session.session_key, status=Dataset.Status.READY)
    source.columns = [{'key': 'c0', 'name': 'Value', 'type': 'string'}]
    source.save()
    url = f'/api/datasets/{source.pk}/replace/'
    values = {'column': 'c0', 'find': 'a', 'replacement': ''}
    assert browser.post(url, {**values, 'column': 'c1'}, format='json').status_code == 400
    assert browser.post(url, {**values, 'find': ''}, format='json').status_code == 400
    response = browser.post(url, values, format='json')
    assert response.status_code == 202 and response.data['kind'] == 'REPLACE_TEXT'
    assert response.data['dataset_id'] != str(source.pk)
    assert Dataset.objects.get(pk=response.data['dataset_id']).parent_id == source.pk


@pytest.mark.spark
def test_real_spark_csv_xlsx_parquet_and_literal_replacement(store, tmp_path):
    # Real local JVM, no eager Celery shortcut for end-to-end HTTP verification (separate smoke script).
    for filename in ['contacts.csv', 'contacts.xlsx']:
        result = dataset(connection=connection('owner'))
        result.source_key = filename
        result.source_format = Path(filename).suffix[1:]
        result.save()
        job = Job.objects.create(owner_session='owner', kind=Job.Kind.IMPORT_DATASET, connection=result.connection, dataset=result)
        def fixture_download(conn, key, target, report):
            import shutil
            shutil.copyfile(SAMPLES / key, target)
            return {'source_bytes': (SAMPLES / key).stat().st_size}
        with patch('processing.services.datasets.download_source', side_effect=fixture_download):
            run_job.run(str(job.pk))
        job.refresh_from_db()
        result.refresh_from_db()
        assert job.status == Job.Status.SUCCESS
        assert result.row_count == 8 and len(result.columns) == 6
        assert result.metadata['partition_count'] == 4
        assert read_page(result, limit=3)['next_cursor'] == 3
        assert read_page(result, after=6, limit=3)['rows'][-1]['row_number'] == 8
        assert list((artifact_directory(result.pk) / 'parquet').glob('part-*.parquet'))
    original = read_page(result)['rows']
    # Literal '.', '$' and backslash exercise Java pattern/replacement escaping.
    output = Dataset.objects.create(owner_session='owner', parent=result, source_key=result.source_key, source_format=result.source_format)
    job = Job.objects.create(owner_session='owner', kind=Job.Kind.REPLACE_TEXT, dataset=output,
        parameters={'column': 'c2', 'find': '.', 'replacement': '$\\'})
    run_job.run(str(job.pk))
    job.refresh_from_db()
    output.refresh_from_db()
    assert job.status == Job.Status.SUCCESS and output.row_count == 8
    changed = read_page(output)['rows']
    for before, after in zip(original, changed):
        value = before['cells'][2]
        assert after['cells'][2] == (value.replace('.', '$\\') if value is not None else None)
    assert read_page(result)['rows'] == original


@pytest.mark.spark
def test_real_spark_header_only_file_creates_empty_preview(store):
    result = dataset(connection=connection('owner'))
    job = Job.objects.create(owner_session='owner', kind=Job.Kind.IMPORT_DATASET, connection=result.connection, dataset=result)
    def fixture_download(conn, key, target, report):
        Path(target).write_text('ID,Name\n')
        return {'source_bytes': 8}
    with patch('processing.services.datasets.download_source', side_effect=fixture_download):
        run_job.run(str(job.pk))
    job.refresh_from_db()
    result.refresh_from_db()
    assert job.status == Job.Status.SUCCESS
    assert result.row_count == 0 and len(result.columns) == 2
    assert read_page(result)['rows'] == []
